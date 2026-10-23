import unittest
from datetime import date, timedelta
from pathlib import Path
from tempfile import TemporaryDirectory
from unittest.mock import patch

from fastapi import HTTPException
from sqlalchemy import create_engine
from sqlalchemy.orm import sessionmaker
from sqlalchemy.pool import StaticPool
from docx import Document as WordDocument

from database import Base
from models import Document, MedicalRecord, Medication, Patient, User
from routes.visit_brief import get_visit_brief
from routes.visit_brief import VisitBriefAnalysisConsent, analyze_visit_brief


class VisitBriefTests(unittest.TestCase):
    def setUp(self):
        self.engine = create_engine(
            "sqlite://",
            connect_args={"check_same_thread": False},
            poolclass=StaticPool,
        )
        Base.metadata.create_all(self.engine)
        self.db = sessionmaker(bind=self.engine, autoflush=False, autocommit=False)()
        self.patient_user = User(
            name="Test Patient",
            email="patient@example.test",
            password="not-a-real-password-hash",
            role="patient",
        )
        self.other_user = User(
            name="Test Doctor",
            email="doctor@example.test",
            password="not-a-real-password-hash",
            role="doctor",
        )
        self.db.add_all([self.patient_user, self.other_user])
        self.db.flush()
        self.patient = Patient(user_id=self.patient_user.id, patient_id="PAT-TEST")
        self.db.add(self.patient)
        self.db.commit()

    def tearDown(self):
        self.db.close()
        self.engine.dispose()

    def test_brief_contains_only_patient_records_and_current_dated_medications(self):
        today = date.today()
        another_patient_user = User(
            name="Another Patient",
            email="another-patient@example.test",
            password="not-a-real-password-hash",
            role="patient",
        )
        self.db.add(another_patient_user)
        self.db.flush()
        another_patient = Patient(
            user_id=another_patient_user.id,
            patient_id="PAT-OTHER",
        )
        self.db.add(another_patient)
        self.db.flush()
        self.db.add_all(
            [
                MedicalRecord(
                    patient_id=self.patient.id,
                    title="Annual visit",
                    diagnosis="Example condition",
                    description="Patient-provided record note",
                    doctor_name="Dr. Example",
                    record_date=today,
                ),
                MedicalRecord(
                    patient_id=another_patient.id,
                    title="Private record belonging to another patient",
                    record_date=today,
                ),
                Medication(
                    patient_id=self.patient.id,
                    name="Current medicine",
                    dosage="1 tablet",
                    frequency="Daily",
                    start_date=today - timedelta(days=2),
                    end_date=today + timedelta(days=5),
                ),
                Medication(
                    patient_id=self.patient.id,
                    name="Past medicine",
                    end_date=today - timedelta(days=1),
                ),
            ]
        )
        self.db.commit()

        brief = get_visit_brief(self.patient_user, self.db)

        self.assertEqual(brief["record_count"], 1)
        self.assertEqual(brief["records"][0]["source"], "Medical record #1")
        self.assertEqual(brief["records"][0]["diagnosis"], "Example condition")
        self.assertEqual(brief["medications_within_date_range_count"], 1)
        current_medication = next(
            medication
            for medication in brief["medications"]
            if medication["name"] == "Current medicine"
        )
        past_medication = next(
            medication
            for medication in brief["medications"]
            if medication["name"] == "Past medicine"
        )
        self.assertTrue(current_medication["within_date_range"])
        self.assertFalse(past_medication["within_date_range"])
        self.assertIn("does not infer what a patient has", brief["notice"])
        self.assertTrue(brief["discussion_questions"])

    def test_non_patient_cannot_generate_visit_brief(self):
        with self.assertRaises(HTTPException) as error:
            get_visit_brief(self.other_user, self.db)

        self.assertEqual(error.exception.status_code, 403)

    def test_analysis_requires_patient_external_processing_consent(self):
        with self.assertRaises(HTTPException) as error:
            import asyncio
            asyncio.run(
                analyze_visit_brief(
                    VisitBriefAnalysisConsent(consent_to_external_processing=False),
                    self.patient_user,
                    self.db,
                )
            )
        self.assertEqual(error.exception.status_code, 400)

    def test_missing_patient_profile_returns_not_found(self):
        unprofiled_patient = User(
            name="New Patient",
            email="new-patient@example.test",
            password="not-a-real-password-hash",
            role="patient",
        )
        self.db.add(unprofiled_patient)
        self.db.commit()

        with self.assertRaises(HTTPException) as error:
            get_visit_brief(unprofiled_patient, self.db)

        self.assertEqual(error.exception.status_code, 404)

    def test_brief_extracts_only_the_signed_in_patients_documents(self):
        with TemporaryDirectory() as upload_directory:
            uploads = Path(upload_directory) / "uploads"
            uploads.mkdir()
            document_path = uploads / "visit.docx"
            document_file = WordDocument()
            document_file.add_paragraph("Documented visit note for this patient.")
            document_file.save(document_path)

            self.db.add(
                Document(
                    patient_id=self.patient.id,
                    title="Visit note",
                    file_name="visit.docx",
                    file_path="uploads/visit.docx",
                )
            )
            another_patient_user = User(
                name="Another Patient",
                email="another-patient@example.test",
                password="not-a-real-password-hash",
                role="patient",
            )
            self.db.add(another_patient_user)
            self.db.flush()
            another_patient = Patient(
                user_id=another_patient_user.id,
                patient_id="PAT-OTHER",
            )
            self.db.add(another_patient)
            self.db.flush()
            self.db.add(
                Document(
                    patient_id=another_patient.id,
                    title="Private note",
                    file_name="private.docx",
                    file_path="private.docx",
                )
            )
            self.db.commit()

            with patch("routes.visit_brief.UPLOAD_DIRECTORY", uploads):
                brief = get_visit_brief(self.patient_user, self.db)

        self.assertEqual(brief["document_count"], 1)
        self.assertEqual(brief["documents_read"], 1)
        self.assertIn("Documented visit note", brief["documents"][0]["text"])
        self.assertIn("Uploaded document #", brief["documents"][0]["source"])

    def test_brief_includes_all_saved_records_and_medications(self):
        today = date.today()
        self.db.add_all(
            [
                MedicalRecord(
                    patient_id=self.patient.id,
                    title=f"Record {index}",
                    record_date=today,
                )
                for index in range(12)
            ]
            + [
                Medication(patient_id=self.patient.id, name=f"Medication {index}")
                for index in range(13)
            ]
        )
        self.db.commit()

        brief = get_visit_brief(self.patient_user, self.db)

        self.assertEqual(len(brief["records"]), 12)
        self.assertEqual(len(brief["medications"]), 13)


if __name__ == "__main__":
    unittest.main()
