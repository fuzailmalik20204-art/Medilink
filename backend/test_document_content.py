import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

import fitz
from docx import Document as WordDocument
from PIL import Image
import pytesseract

from document_content import DocumentContentError, extract_document_content


class DocumentContentTests(unittest.TestCase):
    def setUp(self):
        self.temp_directory = tempfile.TemporaryDirectory()
        self.directory = Path(self.temp_directory.name)

    def tearDown(self):
        self.temp_directory.cleanup()

    def test_extracts_text_from_searchable_pdf(self):
        path = self.directory / "record.pdf"
        pdf = fitz.open()
        page = pdf.new_page()
        page.insert_text((72, 72), "Visit notes: follow-up documented.")
        pdf.save(path)
        pdf.close()

        result = extract_document_content(path, "record.pdf")

        self.assertIn("follow-up documented", result["text"])
        self.assertFalse(result["ocr_used"])
        self.assertEqual(result["status"], "read")

    def test_extracts_paragraphs_and_tables_from_docx(self):
        path = self.directory / "record.docx"
        document = WordDocument()
        document.add_paragraph("Discharge note: follow-up in two weeks.")
        table = document.add_table(rows=1, cols=2)
        table.cell(0, 0).text = "Test"
        table.cell(0, 1).text = "Recorded result"
        document.save(path)

        result = extract_document_content(path, "record.docx")

        self.assertIn("follow-up in two weeks", result["text"])
        self.assertIn("Test | Recorded result", result["text"])

    def test_ocr_extracts_text_from_image(self):
        path = self.directory / "scan.png"
        Image.new("RGB", (4, 4), "white").save(path)

        with patch("document_content.pytesseract.image_to_string", return_value="Scanned record text"):
            result = extract_document_content(path, "scan.png")

        self.assertEqual(result["text"], "Scanned record text")
        self.assertTrue(result["ocr_used"])

    def test_unavailable_ocr_is_reported_clearly(self):
        path = self.directory / "scan.png"
        Image.new("RGB", (4, 4), "white").save(path)

        with patch(
            "document_content.pytesseract.image_to_string",
            side_effect=pytesseract.TesseractNotFoundError(),
        ):
            with self.assertRaisesRegex(DocumentContentError, "OCR is unavailable"):
                extract_document_content(path, "scan.png")

    def test_legacy_doc_is_explicitly_unsupported(self):
        path = self.directory / "record.doc"
        path.write_bytes(b"legacy document")

        with self.assertRaisesRegex(DocumentContentError, "Legacy .doc files"):
            extract_document_content(path, "record.doc")


if __name__ == "__main__":
    unittest.main()
