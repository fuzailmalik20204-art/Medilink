import logging
import zipfile
from pathlib import Path

import fitz
from docx import Document as WordDocument
from docx.opc.exceptions import PackageNotFoundError
from PIL import Image, ImageOps
import pytesseract
from pytesseract import TesseractNotFoundError, TesseractError


logger = logging.getLogger(__name__)
MAX_FILE_BYTES = 15 * 1024 * 1024
MAX_PDF_PAGES = 100
MAX_EXTRACTED_CHARACTERS = 6000
OCR_DPI = 180


class DocumentContentError(Exception):
    pass


def _limit_text(text):
    cleaned = "\n".join(line.strip() for line in text.splitlines() if line.strip())
    if len(cleaned) > MAX_EXTRACTED_CHARACTERS:
        return cleaned[:MAX_EXTRACTED_CHARACTERS].rstrip(), True
    return cleaned, False


def _ocr_image(image):
    try:
        return pytesseract.image_to_string(ImageOps.exif_transpose(image))
    except TesseractNotFoundError as error:
        raise DocumentContentError(
            "OCR is unavailable on this server. Install Tesseract OCR and retry."
        ) from error
    except TesseractError as error:
        raise DocumentContentError("OCR could not read this image.") from error


def _extract_pdf(path):
    pages_text = []
    ocr_used = False
    partial = False

    try:
        with fitz.open(path) as pdf:
            page_count = len(pdf)
            for index, page in enumerate(pdf):
                if index >= MAX_PDF_PAGES:
                    partial = True
                    break

                page_text = page.get_text("text").strip()
                if not page_text:
                    pixmap = page.get_pixmap(dpi=OCR_DPI, alpha=False)
                    image = Image.frombytes(
                        "RGB",
                        (pixmap.width, pixmap.height),
                        pixmap.samples,
                    )
                    page_text = _ocr_image(image).strip()
                    ocr_used = True
                if page_text:
                    pages_text.append(page_text)

            if page_count > MAX_PDF_PAGES:
                partial = True
    except DocumentContentError:
        raise
    except (fitz.FileDataError, fitz.EmptyFileError, RuntimeError, ValueError) as error:
        raise DocumentContentError("The PDF could not be read.") from error

    return "\n\n".join(pages_text), ocr_used, partial


def _extract_docx(path):
    try:
        document = WordDocument(path)
    except (PackageNotFoundError, OSError, ValueError, KeyError, zipfile.BadZipFile) as error:
        logger.warning("Unable to open uploaded DOCX", exc_info=True)
        raise DocumentContentError("The Word document could not be read.") from error

    contents = [paragraph.text for paragraph in document.paragraphs if paragraph.text.strip()]
    for table in document.tables:
        for row in table.rows:
            row_text = " | ".join(cell.text.strip() for cell in row.cells if cell.text.strip())
            if row_text:
                contents.append(row_text)
    return "\n".join(contents), False, False


def extract_document_content(file_path, file_name):
    path = Path(file_path)
    extension = Path(file_name).suffix.lower()
    if not path.is_file():
        raise DocumentContentError(
            "The uploaded file is missing from backend storage. Re-upload it to include its contents."
        )

    try:
        file_size = path.stat().st_size
    except OSError as error:
        logger.warning("Unable to inspect uploaded file", exc_info=True)
        raise DocumentContentError("The uploaded file could not be accessed.") from error

    if file_size > MAX_FILE_BYTES:
        raise DocumentContentError(
            "This file exceeds the 15 MB text-extraction limit and was not read."
        )
    if file_size == 0:
        raise DocumentContentError("This uploaded file is empty.")

    if extension == ".pdf":
        text, ocr_used, partial = _extract_pdf(path)
    elif extension == ".docx":
        text, ocr_used, partial = _extract_docx(path)
    elif extension in {".jpg", ".jpeg", ".png"}:
        try:
            with Image.open(path) as image:
                text = _ocr_image(image)
        except DocumentContentError:
            raise
        except OSError as error:
            raise DocumentContentError("The image could not be read.") from error
        ocr_used, partial = True, False
    elif extension == ".doc":
        raise DocumentContentError(
            "Legacy .doc files are not supported for text extraction. Convert to .docx or PDF and upload again."
        )
    else:
        raise DocumentContentError("This file type is not supported for text extraction.")

    limited_text, text_truncated = _limit_text(text)
    if not limited_text:
        raise DocumentContentError(
            "No readable text was found. The file may be blank, image-only, or use an unsupported scan format."
        )

    return {
        "text": limited_text,
        "ocr_used": ocr_used,
        "truncated": text_truncated or partial,
        "status": "partial" if text_truncated or partial else "read",
        "status_message": _build_status_message(ocr_used, text_truncated, partial),
    }


def _build_status_message(ocr_used, text_truncated, partial):
    if partial:
        return (
            f"Only the first {MAX_PDF_PAGES} PDF pages and/or "
            f"{MAX_EXTRACTED_CHARACTERS} characters were included."
        )
    if text_truncated:
        return f"Text was limited to {MAX_EXTRACTED_CHARACTERS} characters."
    if ocr_used:
        return "Text extracted with OCR; review the original file to verify accuracy."
    return "Text extracted from the original file."
